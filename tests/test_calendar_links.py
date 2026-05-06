from datetime import datetime
from pathlib import Path

from docx import Document

from tools.calendar_links import parse_events_from_docx


def test_parse_events_from_docx_extracts_paragraph_patterns(tmp_path: Path) -> None:
    docx_path = tmp_path / "agenda.docx"
    document = Document()
    document.add_paragraph("12/05/2026 09:00 - 10:30 Reunión inicial")
    document.add_paragraph("12 de mayo de 2026, 11:00 a 12:30 - Taller")
    document.save(docx_path)

    events = parse_events_from_docx(str(docx_path))

    assert [event.title for event in events] == ["Reunión inicial", "Taller"]
    assert events[0].start_datetime == datetime(2026, 5, 12, 9, 0)
    assert events[0].end_datetime == datetime(2026, 5, 12, 10, 30)
    assert events[1].start_datetime == datetime(2026, 5, 12, 11, 0)
    assert events[1].end_datetime == datetime(2026, 5, 12, 12, 30)
    assert all(event.timezone == "UTC" for event in events)


def test_parse_events_from_docx_extracts_table_rows(tmp_path: Path) -> None:
    docx_path = tmp_path / "agenda_tabla.docx"
    document = Document()
    table = document.add_table(rows=1, cols=4)
    for cell, value in zip(table.rows[0].cells, ["Fecha", "Hora", "Actividad", "Lugar"]):
        cell.text = value
    row = table.add_row().cells
    row[0].text = "12/05/2026"
    row[1].text = "09:00 - 10:30"
    row[2].text = "Reunión inicial"
    row[3].text = "Sala 1"
    document.save(docx_path)

    events = parse_events_from_docx(str(docx_path))

    assert len(events) == 1
    assert events[0].title == "Reunión inicial"
    assert events[0].description == "Reunión inicial"
    assert events[0].location == "Sala 1"
    assert events[0].start_datetime == datetime(2026, 5, 12, 9, 0)
    assert events[0].end_datetime == datetime(2026, 5, 12, 10, 30)
